"""Fill a Word template's content controls and attach a Word comment per citation.

Usage:
    from docfill.fill import fill
    fill("credit_paper_template.docx", {
        "exec_summary": {"kind": "text", "text": "...", "citations": [{"source": "valuation_report.pdf", "page": 2, "quote": "..."}]},
        "terms_table": {"kind": "table", "rows": [["Term", "Detail"], ["Limit", "A$185.0M"]], "citations": [...]},
        "site_plan_image": {"kind": "image", "image_path": "site.png", "caption": "Site plan", "citations": [...]},
        "lvr": {"kind": "metric", "text": "58.2%", "citations": [{"source": "mv_deal_metrics", "measure": "LVR"}]},
    }, "out.docx")

Each content control keeps its tag, so a filled document can be filled again or inspected.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import ooxml

AUTHOR = "IC paper agent"
INITIALS = "AI"


@dataclass
class FillReport:
    filled: list[str] = field(default_factory=list)
    missing_in_template: list[str] = field(default_factory=list)
    unfilled: list[str] = field(default_factory=list)
    comments: int = 0


def list_slots(template_path: str | Path) -> list[dict[str, str]]:
    """Return [{tag, alias, kind}] for every content control in document order."""
    doc = Document(str(template_path))
    out = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        tag = pr.find(qn("w:tag"))
        alias = pr.find(qn("w:alias"))
        if tag is None:
            continue
        if pr.find(qn("w:picture")) is not None:
            kind = "image"
        elif pr.find(qn("w:text")) is not None or sdt.getparent().tag == qn("w:p"):
            kind = "metric"
        else:
            kind = "block"
        out.append(dict(tag=tag.get(qn("w:val")), alias=alias.get(qn("w:val")) if alias is not None else "", kind=kind))
    return out


def format_citation(c: dict[str, Any]) -> str:
    if c.get("label"):
        return c["label"]
    parts = []
    src = c.get("source", "")
    if c.get("type") == "draft":
        return "Agent draft, not from a source: needs the analyst's decision."
    ref = f"[{c['ref']}] " if c.get("ref") else ""
    if c.get("measure") or c.get("type") == "metric":
        what = "field" if c.get("dimension") else "measure"
        value = f" = {c['value']}" if c.get("value") else ""
        name = c.get("view_label") or src.split(".")[-1]
        parts.append(f"{ref}{c['measure']}{value} · {name} ({what}, Unity Catalog metric view)" if c.get("measure")
                     else f"{ref}{name} (Unity Catalog metric view)")
        if c.get("query"):
            parts.append(f"Query: {c['query']}")
        if c.get("as_of"):
            parts.append(f"As of: {c['as_of']}")
    else:
        loc = f"{ref}{c.get('title') or src}" + (f" · {c['issuer']}" if c.get("issuer") else "")
        if c.get("page"):
            loc += f", page {c['page']}"
        if c.get("section"):
            loc += f", {c['section']}"
        parts.append(loc)
        if c.get("quote"):
            parts.append(f"“{c['quote']}”")
    if c.get("url"):
        parts.append(c["url"])
    return "\n".join(parts)


def _clear_placeholder(sdt) -> None:
    pr = sdt.find(qn("w:sdtPr"))
    for flag in pr.findall(qn("w:showingPlcHdr")):
        pr.remove(flag)


def _content(sdt):
    return sdt.find(qn("w:sdtContent"))


def _first_pstyle(content) -> str | None:
    p = content.find(qn("w:p"))
    if p is None:
        return None
    ps = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else None


def _add_comments(doc, runs: list[Run], citations: list[dict[str, Any]], report: FillReport) -> None:
    if not runs or not citations:
        return
    for c in citations:
        doc.add_comment(runs if len(runs) > 1 else runs[0], text=format_citation(c), author=AUTHOR, initials=INITIALS)
        report.comments += 1


def _new_paragraph(doc, text: str, style: str | None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    run = p.add_run(text)
    p._p.getparent().remove(p._p)
    return p, run


def _fill_text(doc, sdt, value: dict[str, Any], report: FillReport) -> None:
    content = _content(sdt)
    style_id = _first_pstyle(content)
    style = None
    if style_id:
        style = next((s.name for s in doc.styles if getattr(s, "style_id", None) == style_id), None)
    for child in list(content):
        content.remove(child)

    items: list[dict[str, Any]] = []
    if value.get("paragraphs"):
        items = [p if isinstance(p, dict) else {"text": p} for p in value["paragraphs"]]
    elif value.get("bullets"):
        items = [b if isinstance(b, dict) else {"text": b} for b in value["bullets"]]
        for it in items:
            it.setdefault("style", "List Bullet")
    else:
        text = value.get("text", "")
        texts = text if isinstance(text, list) else [t for t in str(text).split("\n\n") if t.strip()]
        items = [{"text": t} for t in texts]

    all_runs: list[Run] = []
    for it in items:
        # a paragraph is one run, or several when it carries formatting: runs=[{"text", "bold", "italic"}]
        runs_in = it.get("runs") or [{"text": it["text"]}]
        p, _ = _new_paragraph(doc, runs_in[0]["text"], it.get("style") or style)
        content.append(p._p)
        para = Paragraph(p._p, doc._body)
        for extra in runs_in[1:]:
            para.add_run(extra["text"])
        runs = list(para.runs)
        for r, spec in zip(runs, runs_in):
            r.bold = True if spec.get("bold") else None
            r.italic = True if spec.get("italic") else None
        all_runs.extend(runs)
        _add_comments(doc, runs, it.get("citations", []), report)
    _add_comments(doc, all_runs, value.get("citations", []), report)


def _style_table(table, header: bool, col_widths_cm: list[float] | None) -> None:
    ooxml.table_borders(table, color="C9D4E3")
    ooxml.cell_margins(table, 50, 50, 80, 80)
    for r_i, row in enumerate(table.rows):
        for c_i, cell in enumerate(row.cells):
            if col_widths_cm and c_i < len(col_widths_cm):
                cell.width = Cm(col_widths_cm[c_i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if header and r_i == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if header and r_i == 0:
                ooxml.shade_cell(cell, ooxml.NAVY)
            elif r_i % 2 == 0:
                ooxml.shade_cell(cell, "F4F8FC")
    if header:
        ooxml.repeat_header(table.rows[0])


def _fill_table(doc, sdt, value: dict[str, Any], report: FillReport) -> None:
    rows = value["rows"]
    content = _content(sdt)
    for child in list(content):
        content.remove(child)
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    for r_i, row in enumerate(rows):
        for c_i in range(n_cols):
            table.cell(r_i, c_i).paragraphs[0].add_run("" if c_i >= len(row) else str(row[c_i]))
    _style_table(table, value.get("header", True), value.get("col_widths_cm"))
    tbl = table._tbl
    tbl.getparent().remove(tbl)
    content.append(tbl)
    anchor: list[Run] = []
    if value.get("caption"):
        p, _ = _new_paragraph(doc, value["caption"], "Caption")
        content.append(p._p)
        anchor = [Paragraph(p._p, doc._body).runs[0]]
    else:
        first = table.cell(0, 0).paragraphs[0].runs
        anchor = [first[0]] if first else []
    _add_comments(doc, anchor, value.get("citations", []), report)
    # per-row citations: {"row_citations": {row_index: [citations]}}
    for r_i, cits in (value.get("row_citations") or {}).items():
        runs = table.cell(int(r_i), 0).paragraphs[0].runs
        _add_comments(doc, runs[:1], cits, report)


def _fill_image(doc, sdt, value: dict[str, Any], report: FillReport) -> None:
    content = _content(sdt)
    for child in list(content):
        content.remove(child)
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(str(value["image_path"]), width=Cm(value.get("width_cm", 16)))
    p._p.getparent().remove(p._p)
    content.append(p._p)
    anchor = [Paragraph(p._p, doc._body).runs[0]]
    if value.get("caption"):
        cap, _ = _new_paragraph(doc, value["caption"], "Caption")
        sdt.addnext(cap._p)  # caption sits outside the picture control
        anchor = [Paragraph(cap._p, doc._body).runs[0]]
    _add_comments(doc, anchor, value.get("citations", []), report)


def _fill_metric(doc, sdt, value: dict[str, Any], report: FillReport) -> None:
    pr = sdt.find(qn("w:sdtPr"))
    for t in pr.findall(qn("w:text")):
        pr.remove(t)  # rich text inline control so a comment range can sit inside it
    content = _content(sdt)
    for child in list(content):
        content.remove(child)
    r = ooxml._el("w:r")
    rpr = ooxml._el("w:rPr")
    rpr.append(ooxml._el("w:b"))
    r.append(rpr)
    t = ooxml._el("w:t")
    t.text = str(value.get("text", ""))
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    content.append(r)
    para_el = sdt.getparent()
    _add_comments(doc, [Run(r, Paragraph(para_el, doc._body))], value.get("citations", []), report)


FILLERS = {"text": _fill_text, "table": _fill_table, "image": _fill_image, "metric": _fill_metric}


def fill(template_path: str | Path, values: dict[str, dict[str, Any]], out_path: str | Path) -> FillReport:
    doc = Document(str(template_path))
    report = FillReport()
    by_tag = {}
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag = sdt.find(qn("w:sdtPr") + "/" + qn("w:tag"))
        if tag is not None:
            by_tag[tag.get(qn("w:val"))] = sdt
    for tag, value in values.items():
        sdt = by_tag.get(tag)
        if sdt is None:
            report.missing_in_template.append(tag)
            continue
        kind = value.get("kind", "text")
        FILLERS[kind](doc, sdt, value, report)
        _clear_placeholder(sdt)
        report.filled.append(tag)
    report.unfilled = [t for t in by_tag if t not in values]
    doc.save(str(out_path))
    return report
